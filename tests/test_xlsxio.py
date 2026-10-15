"""Browser-free tests for the xlsx workbook (plans/xlsx-workbook-format.md)."""

from __future__ import annotations

import stat
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from concurrer.csvio import ExpenseRecord, FieldCell
from concurrer.type_form_map import TYPE_FORM_FIELDS
from concurrer.type_map import TYPE_FIELDS
from concurrer.xlsxio import (
    CUSTOMIZATION,
    DETAILS,
    EXPENSES,
    LISTS,
    READ_ONLY,
    WorkbookData,
    customization_rows,
    fit_width,
    plan_columns,
    verify_workbook,
    write_workbook,
    workbook_issues,
)

REPORT = "506EC055FE79452BB943"


def cell(name, label, value="", locked=False, required=False):
    return FieldCell(name, label, value, locked, required)


def gt():
    return ExpenseRecord("AAAA1111", REPORT, 1, [
        cell("expenseType", "Expense Type", "Ground Transportation", required=True),
        cell("transactionDate", "Transaction Date", "09/21/2026", locked=True),
        cell("vendorName", "Vendor Name", "INTERNATIONAL TRANSACTION", locked=True),
        cell("custom12", "Business Purpose", "Student Life", required=True),
        cell("custom20", "Type of Ground Transportation", "Taxi", required=True),
        cell("transactionCurrencyName", "Currency", "US, Dollar (USD)", locked=True),
        cell("comment", "Comment", 'has, commas; and "quotes"'),
    ])


def air():
    return ExpenseRecord("BBBB2222", REPORT, 2, [
        cell("expenseType", "Expense Type", "Airfare", required=True),
        cell("custom20", "Type of Class", "", required=True),
        cell("vendor", "Vendor Name", "United Airlines"),
        cell("location", "Destination", "Houston, Texas"),
    ])


def data():
    return WorkbookData(
        REPORT, [gt(), air()], {"Report Name": "CCED Summer 2026"},
        {("custom20", "type of class"): ["01. Economy", "02. Business Class"],
         ("custom12", "business purpose"): ["Student Life", "Group Travel"]},
        {"custom21": 10},
    )


class XlsxTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "r.xlsx"

    def tearDown(self):
        self.tmp.cleanup()

    def test_custom20_meanings_merge_into_variant(self):
        headers = [c.header for c in plan_columns([gt(), air()])]
        self.assertIn("Variant", headers)
        for gone in ("Type of Ground Transportation", "Type of Class", "custom20"):
            self.assertNotIn(gone, headers)

    def test_uf_order_and_full_schema_columns(self):
        headers = [c.header for c in plan_columns([gt()])]
        # The operator's chosen leading order remains stable, but columns for
        # types absent from this report must exist for a later type change.
        self.assertEqual(headers[:2], ["Date", "Expense Type"])
        self.assertLess(headers.index("Location"), headers.index("Expense ID"))
        self.assertLess(headers.index("Desc./Justification"), headers.index("Expense ID"))
        self.assertLess(headers.index("Vehicle Govt #"), headers.index("Expense ID"))
        self.assertEqual(headers[-1], "Expense ID")
        self.assertEqual(len(headers), len(set(headers)))
        for banned in ("report_id", "__row", "__locked", "__required"):
            self.assertNotIn(banned, headers)

    def test_round_trip_and_permissions(self):
        write_workbook(self.path, data())
        self.assertEqual(verify_workbook(self.path, data()), [])
        self.assertEqual(stat.S_IMODE(self.path.stat().st_mode), 0o600)

    def test_sheets_banner_and_report_id(self):
        write_workbook(self.path, data())
        wb = load_workbook(self.path)
        self.assertEqual(wb.sheetnames,
                         [EXPENSES, "Meals & Per Diem", DETAILS, CUSTOMIZATION,
                          LISTS, "Variants", "Orig"])
        self.assertEqual(wb[LISTS].sheet_state, "hidden")
        self.assertEqual(wb["Variants"].sheet_state, "hidden")
        for name in (DETAILS, CUSTOMIZATION):
            a1 = wb[name]["A1"]
            self.assertEqual(a1.value, READ_ONLY)
            self.assertTrue(a1.font.bold and a1.font.italic)
        facts = {r[0]: r[1] for r in wb[DETAILS].iter_rows(min_row=4, values_only=True)}
        self.assertEqual(facts["Report ID"], REPORT)

    def test_locked_cells_protected_and_editable_cells_open(self):
        write_workbook(self.path, data())
        ws = load_workbook(self.path)[EXPENSES]
        self.assertTrue(ws.protection.sheet)
        headers = [c.value for c in ws[1]]
        date = ws.cell(row=2, column=headers.index("Date") + 1)
        comment = ws.cell(row=2, column=headers.index("Comment") + 1)
        self.assertTrue(date.protection.locked)
        self.assertFalse(comment.protection.locked)
        # Blocked cells are black with bold white type (operator 2026-10-13).
        self.assertEqual(date.fill.fgColor.rgb, "00000000")
        self.assertEqual(date.font.color.rgb, "00FFFFFF")
        self.assertTrue(date.font.bold)
        # Type-specific cells stay editable: a type change can make them apply.
        days = ws.cell(row=2, column=headers.index("Days/Nights") + 1)
        self.assertFalse(days.protection.locked)

    def test_dropdowns_and_applicability_rules(self):
        write_workbook(self.path, data())
        ws = load_workbook(self.path)[EXPENSES]
        formulas = [dv.formula1 for dv in ws.data_validations.dataValidation]
        self.assertTrue(any(f.startswith("Lists!") for f in formulas))
        rules = [r.formula[0] for rng in ws.conditional_formatting
                 for r in rng.rules]
        self.assertTrue(any("COUNTIFS('Backend Data'!$A:$A" in r for r in rules))
        # Full state system: required-empty pink, required-filled green,
        # optional grey, not-applicable grey/red, list/whole-number reds.
        fills = [r.dxf.fill.fgColor.rgb for rng in ws.conditional_formatting
                 for r in rng.rules if r.dxf and r.dxf.fill and r.dxf.fill.fgColor]
        for expected in ("00F8D7DA", "00C6EFCE", "00F2F2F2", "00D9D9D9", "00FF9999"):
            self.assertIn(expected, fills)
        self.assertTrue(any('="Yes",FALSE)' in r for r in rules))
        self.assertTrue(any("COUNTIF(Lists!" in r for r in rules))

    def test_workbook_issue_mirror_matches_red_and_edited_pink_rules(self):
        write_workbook(self.path, data())
        self.assertEqual(workbook_issues(self.path), [])
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        headers = [c.value for c in ws[1]]

        # An edited GT row has an option outside its Variant list: bright red.
        ws.cell(2, headers.index("Variant") + 1).value = "Plane"
        wb.save(self.path)
        self.assertIn("off-list", {issue.kind for issue in workbook_issues(self.path)})

        # Required-empty is pink, and blocks only because this row was edited.
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        ws.cell(2, headers.index("Variant") + 1).value = None
        wb.save(self.path)
        self.assertIn("required-empty", {issue.kind for issue in workbook_issues(self.path)})

        # A value in the target-inapplicable Days/Nights custom column is red.
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        ws.cell(2, headers.index("Days/Nights") + 1).value = "3"
        wb.save(self.path)
        self.assertIn("not-applicable", {issue.kind for issue in workbook_issues(self.path)})

        # An edited locked card field is red even if it otherwise looks valid.
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        ws.cell(2, headers.index("Date") + 1).value = "10/01/2026"
        wb.save(self.path)
        self.assertIn("locked-deviation", {issue.kind for issue in workbook_issues(self.path)})

    def test_unverified_applicability_is_yellow_not_optional_grey(self):
        # Operator 2026-10-14: light grey read as "allowed". A standard field
        # with no Backend Data row for the row's type is unknown, not optional.
        write_workbook(self.path, data())
        ws = load_workbook(self.path)[EXPENSES]
        headers = [c.value for c in ws[1]]
        letter = get_column_letter(headers.index("Location") + 1)
        rules = [r for rng in ws.conditional_formatting
                 if str(rng.sqref).startswith(f"{letter}2:") for r in rng.rules]
        fills = [r.dxf.fill.fgColor.rgb for r in rules]
        self.assertIn("00FFFF99", fills)
        # Yellow is evaluated before (and stops) the optional grey.
        self.assertLess(fills.index("00FFFF99"), fills.index("00F2F2F2"))
        self.assertTrue(rules[fills.index("00FFFF99")].stopIfTrue)

    def test_new_value_in_an_unverified_cell_is_red_and_blocks(self):
        # Airfare -> Phone/Fax/Internet: the latter's blank form has no
        # Desc./Justification, so a new value there has nowhere to go.
        write_workbook(self.path, data())
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        headers = [c.value for c in ws[1]]
        ws.cell(3, headers.index("Expense Type") + 1).value = "Phone/Fax/Internet"
        ws.cell(3, headers.index("Desc./Justification") + 1).value = "Flight Wifi"
        wb.save(self.path)
        flagged = {(i.header, i.kind) for i in workbook_issues(self.path)}
        self.assertIn(("Desc./Justification", "unverified-new"), flagged)
        # The unchanged exported Location on that row stays yellow, not red.
        self.assertNotIn(("Location", "unverified-new"), flagged)
        letter = get_column_letter(headers.index("Desc./Justification") + 1)
        rules = [r for rng in ws.conditional_formatting
                 if str(rng.sqref).startswith(f"{letter}2:") for r in rng.rules]
        fills = [r.dxf.fill.fgColor.rgb for r in rules]
        red_new = next(i for i, r in enumerate(rules)
                       if "Orig!" in r.formula[0] and r.dxf.fill.fgColor.rgb == "00FF9999")
        self.assertLess(red_new, fills.index("00FFFF99"))

    def test_vendor_column_gets_extra_width(self):
        write_workbook(self.path, data())
        ws = load_workbook(self.path)[EXPENSES]
        headers = [c.value for c in ws[1]]
        letter = get_column_letter(headers.index("vendor") + 1)
        values = [ws.cell(r, headers.index("vendor") + 1).value for r in range(2, ws.max_row + 1)]
        self.assertGreater(ws.column_dimensions[letter].width, fit_width("vendor", values))

    def test_untouched_off_list_value_does_not_block_the_batch(self):
        # A freshly exported "Undefined" type is Concur's own state: it is
        # never written, so it must not block an unrelated edit elsewhere.
        undef = ExpenseRecord("DDDD4444", REPORT, 3, [
            cell("expenseType", "Expense Type", "Undefined", required=True),
        ])
        write_workbook(self.path, WorkbookData(REPORT, [gt(), undef], {}, {}, {}))
        self.assertNotIn("off-list", {i.kind for i in workbook_issues(self.path)})
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        headers = [c.value for c in ws[1]]
        ws.cell(2, headers.index("Comment") + 1).value = "edited"
        wb.save(self.path)
        self.assertNotIn("off-list", {i.kind for i in workbook_issues(self.path)})
        # Typing an off-list value is still refused.
        ws.cell(2, headers.index("Variant") + 1).value = "Plane"
        wb.save(self.path)
        self.assertIn("off-list", {i.kind for i in workbook_issues(self.path)})

    def test_exchange_fields_do_not_apply_to_usd_rows(self):
        # Concur shows Exchange Rate / Amount in USD only on foreign-currency
        # expenses, so live evidence from a JPY row of the same type must not
        # make them required on a USD row (operator 2026-10-13).
        jpy = ExpenseRecord("CCCC3333", REPORT, 3, [
            cell("expenseType", "Expense Type", "Ground Transportation", required=True),
            cell("transactionCurrencyName", "Currency", "Japan, Yen (JPY)", locked=True),
            cell("exchangeRate", "Exchange Rate", "0.0064", locked=True, required=True),
            cell("postedAmount", "Amount in USD", "7.05", locked=True, required=True),
        ])
        write_workbook(self.path, WorkbookData(REPORT, [gt(), jpy], {}, {}, {}))
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        headers = [c.value for c in ws[1]]
        # Edit the USD row so the edited-row pink rule is active for it.
        ws.cell(2, headers.index("Comment") + 1).value = "edited"
        wb.save(self.path)
        flagged = {(i.header, i.kind) for i in workbook_issues(self.path)}
        for header in ("Exchange Rate", "Amount in USD"):
            self.assertNotIn((header, "required-empty"), flagged)
            letter = get_column_letter(headers.index(header) + 1)
            rules = [r.formula[0] for rng in ws.conditional_formatting
                     if str(rng.sqref).startswith(f"{letter}2:")
                     for r in rng.rules]
            self.assertTrue(any('SEARCH("(USD)"' in f for f in rules))

    def test_request_payment_and_currency_are_constrained_dropdowns(self):
        # Payment/Currency are captured fixed Concur lists; Request is the
        # exact report-local values already observed on applicable expenses.
        requested = ExpenseRecord("RRRR0000", REPORT, 3, [
            cell("expenseType", "Expense Type", "Ground Transportation"),
            cell("travelRequestExpenseId", "Request",
                 "05/04/2026, $800.00 - CCED Summer 2026"),
        ])
        write_workbook(self.path, WorkbookData(REPORT, [gt(), requested]))
        wb = load_workbook(self.path)
        ws, lists = wb[EXPENSES], wb[LISTS]
        headers = [c.value for c in ws[1]]
        by_col = {}
        for dv in ws.data_validations.dataValidation:
            for rng in dv.sqref.ranges:
                by_col.setdefault(rng.min_col, dv)
        for header in ("Request", "Payment Type", "Currency"):
            self.assertEqual(by_col[headers.index(header) + 1].type, "list")
        list_headers = [c.value for c in lists[1]]
        payment_col = list_headers.index("Payment Type") + 1
        currency_col = list_headers.index("Currency") + 1
        self.assertEqual(lists.cell(2, payment_col).value, "Pending Transaction")
        self.assertEqual(lists.cell(3, payment_col).value, "Personal Funds")
        self.assertEqual(lists.cell(2, currency_col).value,
                         "Afghanistan, Afghani (AFN)")

    def test_locked_card_value_is_not_recoloured_as_target_type_good(self):
        write_workbook(self.path, data())
        ws = load_workbook(self.path)[EXPENSES]
        headers = [c.value for c in ws[1]]
        date = ws.cell(row=2, column=headers.index("Date") + 1)
        self.assertTrue(date.protection.locked)
        formulas = [r.formula[0] for rng in ws.conditional_formatting
                    for r in rng.rules if f"{date.column_letter}$1" in r.formula[0]]
        self.assertTrue(any(f'CELL("protect",{date.column_letter}2)=0' in f
                            for f in formulas))

    def test_customization_covers_every_full_form_type(self):
        types = {row[0] for row in customization_rows([gt(), air()], {})}
        self.assertEqual(types, set(TYPE_FORM_FIELDS))
        rows = {(r[0], r[1]): r for r in customization_rows([gt(), air()], {})}
        self.assertEqual(rows[("Airfare", "Variant")][2:5],
                         ("Type of Class", "custom20", "Yes"))
        # No current Airfare record is needed to know Phone/Fax/Internet
        # still has Date; it is a positive fact from the full type sweep.
        self.assertEqual(rows[("Phone/Fax/Internet", "Date")][2:5],
                         ("Transaction Date", "transactionDate", "Yes"))

    def test_width_fits_longest_but_ignores_outlier(self):
        values = ["Taxi", "Bus", "Shuttle", "Ferry", "Taxi", "Bus", "Ferry",
                  "x" * 200]
        self.assertEqual(fit_width("Type", values), 8 + 2)  # ceil(7*1.05)+2
        self.assertEqual(fit_width("Ground Transportation", ["a"]), 15 + 2)
        self.assertEqual(fit_width("H", []), 8)

    def test_vendor_routes_locked_card_and_unlocked_directory_fields(self):
        data = WorkbookData(REPORT, [gt(), air()])
        headers = [c.header for c in plan_columns(data.records)]
        self.assertEqual(headers.count("vendor"), 1)
        self.assertNotIn("Vendor Name (list)", headers)
        write_workbook(self.path, data)
        self.assertEqual(verify_workbook(self.path, data), [])
        ws = load_workbook(self.path)[EXPENSES]
        headers = [c.value for c in ws[1]]
        vendor = headers.index("vendor") + 1
        locked = ws.cell(row=2, column=vendor)
        unlocked = ws.cell(row=3, column=vendor)
        self.assertEqual((locked.value, unlocked.value),
                         ("INTERNATIONAL TRANSACTION", "United Airlines"))
        self.assertTrue(locked.protection.locked)
        self.assertFalse(unlocked.protection.locked)

    def test_read_back_finds_columns_by_header_and_rejects_duplicate_ids(self):
        from openpyxl import load_workbook as lw
        from concurrer.xlsxio import read_workbook
        write_workbook(self.path, data())
        wb = lw(self.path)
        ws = wb[EXPENSES]
        ws.move_range("A1:A3", cols=ws.max_column)  # move Date to the end
        ws.delete_cols(1)
        wb.save(self.path)
        _, rows = read_workbook(self.path)
        self.assertEqual(rows["AAAA1111"][("transactionDate", "")], "09/21/2026")
        wb = lw(self.path)
        ws = wb[EXPENSES]
        id_col = [c.value for c in ws[1]].index("Expense ID") + 1
        ws.cell(row=3, column=id_col, value="AAAA1111")
        wb.save(self.path)
        with self.assertRaisesRegex(ValueError, "duplicate Expense ID"):
            read_workbook(self.path)

    def test_days_nights_merges_by_type_and_validates_whole_numbers(self):
        lodging = ExpenseRecord("CCCC3333", REPORT, 3, [
            cell("expenseType", "Expense Type", "Lodging - Group", required=True),
            cell("custom24", "Number of Nights", "4", required=True),
            cell("custom28", "Number of Rooms", "2", required=True),
        ])
        laundry = ExpenseRecord("DDDD4444", REPORT, 4, [
            cell("expenseType", "Expense Type", "Laundry Services", required=True),
            cell("custom28", "Number of consecutive days of travel", "6"),
        ])
        d = WorkbookData(REPORT, [gt(), lodging, laundry])
        cols = {c.header: c for c in plan_columns(d.records)}
        self.assertIn("Days/Nights", cols)
        self.assertIn("# Rooms", cols)  # not a day/night count
        write_workbook(self.path, d)
        self.assertEqual(verify_workbook(self.path, d), [])
        ws = load_workbook(self.path)[EXPENSES]
        headers = [c.value for c in ws[1]]
        col = headers.index("Days/Nights") + 1
        self.assertEqual([ws.cell(row=r, column=col).value for r in (2, 3, 4)],
                         [None, "4", "6"])
        self.assertTrue(any(dv.type == "whole"
                            for dv in ws.data_validations.dataValidation))

    def test_days_nights_column_present_even_when_unused(self):
        self.assertIn("Days/Nights", [c.header for c in plan_columns([gt()])])

    def test_customization_routes_day_night_fields_to_merged_header(self):
        rows = {(t, h, f) for t, h, _l, f, *_ in customization_rows([gt(), air()], {})}
        self.assertIn(("Lodging - Group", "Days/Nights", "custom24"), rows)
        self.assertIn(("Lodging - Group", "# Rooms", "custom28"), rows)
        opts = {(t, h): o for t, h, _l, _f, _r, o, _e in customization_rows([gt(), air()], {})}
        self.assertIn("Van", opts[("Car Rental", "Variant")])
        # Same label, different lists: options must come per type.
        self.assertNotIn(("Travel Meals - Domestic",
                          "Departure Time (first day of travel)"), opts)

    def test_ambiguous_label_gets_no_fixed_dropdown(self):
        from concurrer.xlsxio import OPTIONS_BY_LABEL
        self.assertIsNone(OPTIONS_BY_LABEL[("custom17", "departure time (first day of travel)")])

    def test_per_diem_fields_live_on_meals_sheet_not_expenses(self):
        meals = ExpenseRecord("EEEE5555", REPORT, 5, [
            cell("expenseType", "Expense Type", "Travel Meals - Domestic"),
            cell("custom12", "Business Purpose", "Student Life"),
            cell("custom17", "Departure Time (first day of travel)", "5. Whole Day"),
            cell("custom24", "Number of Breakfasts", "1"),
            cell("custom26", "Number of Lunches", "1"),
        ])
        d = WorkbookData(REPORT, [gt(), meals])
        headers = [c.header for c in plan_columns(d.records)]
        for gone in ("Departure Time (first day of travel)",
                     "Number of Breakfasts", "Number of Lunches"):
            self.assertNotIn(gone, headers)
        write_workbook(self.path, d)
        self.assertEqual(verify_workbook(self.path, d), [])
        from concurrer.xlsxio import read_workbook
        _, rows = read_workbook(self.path)
        self.assertNotIn("EEEE5555", rows)  # meal rows never feed write-back
        rows_c = {(t, h) for t, h, *_ in customization_rows([gt(), air()], {})}
        self.assertNotIn(("Travel Meals - Domestic", "Number of Breakfasts"), rows_c)
        self.assertNotIn(("All Inclusive Per Diem (Meals/Lodging)", "Days/Nights"), rows_c)

    def test_with_others_merges_custom29_with_yes_no_dropdown(self):
        gt_row = ExpenseRecord("FFFF6666", REPORT, 6, [
            cell("expenseType", "Expense Type", "Ground Transportation"),
            cell("custom29", "Were other travelers with you?", "No"),
        ])
        lodging = ExpenseRecord("GGGG7777", REPORT, 7, [
            cell("expenseType", "Expense Type", "Lodging"),
            cell("custom29", "Are you sharing a room?", "Yes"),
        ])
        d = WorkbookData(REPORT, [gt_row, lodging])
        headers = [c.header for c in plan_columns(d.records)]
        self.assertIn("Shared/Accompanied?", headers)
        self.assertNotIn("Were other travelers with you?", headers)
        write_workbook(self.path, d)
        self.assertEqual(verify_workbook(self.path, d), [])
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        col = [c.value for c in ws[1]].index("Shared/Accompanied?") + 1
        self.assertEqual([ws.cell(row=r, column=col).value for r in (2, 3)],
                         ["No", "Yes"])
        letter = ws.cell(row=1, column=col).column_letter
        dv = next(dv for dv in ws.data_validations.dataValidation
                  if f"{letter}2" in str(dv.sqref))
        self.assertEqual(dv.type, "list")

    def test_attendee_and_room_counts_renamed_and_whole_number(self):
        party = ExpenseRecord("HHHH8888", REPORT, 8, [
            cell("expenseType", "Expense Type", "Entertainment Expense"),
            cell("custom28", "Total Number of Attendees", "12"),
        ])
        group = ExpenseRecord("IIII9999", REPORT, 9, [
            cell("expenseType", "Expense Type", "Lodging - Group"),
            cell("custom28", "Number of Rooms", "3"),
            cell("custom24", "Number of Nights", "2"),
        ])
        headers = [c.header for c in plan_columns([party, group])]
        self.assertIn("# Attendees", headers)
        self.assertIn("# Rooms", headers)
        rows = {(t, h) for t, h, *_ in customization_rows([gt(), air()], {})}
        self.assertIn(("Food & Beverages - Lunch", "# Attendees"), rows)
        self.assertIn(("Lodging - Group", "# Rooms"), rows)
        self.assertIn(("Fuel - State Vehicle", "Vehicle Govt #"), rows)
        labels = {(r[0], r[1]): r[2] for r in customization_rows([gt(), air()], {})}
        self.assertEqual(labels[("Fuel - State Vehicle", "Vehicle Govt #")],
                         "State Vehicle/Decal/Vessel Number")

    def test_variant_dependent_dropdown_reads_type_column(self):
        write_workbook(self.path, data())
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        headers = [c.value for c in ws[1]]
        col = ws.cell(row=1, column=headers.index("Variant") + 1).column_letter
        type_col = ws.cell(row=1, column=headers.index("Expense Type") + 1).column_letter
        dv = next(dv for dv in ws.data_validations.dataValidation
                  if f"{col}2" in str(dv.sqref))
        self.assertTrue(dv.formula1.startswith("OFFSET(Variants!$A$2,0,MATCH("))
        self.assertIn(f"${type_col}2", dv.formula1)
        self.assertEqual(dv.formula1.count("("), dv.formula1.count(")"))
        variants = wb["Variants"]
        heads = [c.value for c in variants[1]]
        rental = heads.index("Rental of Space - External") + 1
        self.assertEqual([variants.cell(row=r, column=rental).value for r in (2, 3)],
                         ["12 months or less", "More than 12 months"])
        # GT row keeps its value, now under Variant.
        self.assertEqual(ws.cell(row=2, column=headers.index("Variant") + 1).value, "Taxi")


    def test_absent_request_is_unlocked_for_a_type_change(self):
        # An Undefined row has no Request; switching it to Ground
        # Transportation makes Request required. A pink cell cannot be locked.
        undef = ExpenseRecord("JJJJ0000", REPORT, 10, [
            cell("expenseType", "Expense Type", "Undefined"),
        ])
        write_workbook(self.path, WorkbookData(REPORT, [undef]))
        ws = load_workbook(self.path)[EXPENSES]
        headers = [c.value for c in ws[1]]
        request = ws.cell(row=2, column=headers.index("Request") + 1)
        self.assertFalse(request.protection.locked)

    def test_sparse_standard_fields_use_live_lookup_not_static_grey(self):
        write_workbook(self.path, data())
        wb = load_workbook(self.path)
        ws = wb[EXPENSES]
        headers = [c.value for c in ws[1]]
        loc = headers.index("Location") + 1
        # Ground Transportation row has no Location: white statically, grey by rule.
        gt_cell = ws.cell(row=2, column=loc)
        self.assertIsNone(gt_cell.value)
        self.assertNotEqual(
            (gt_cell.fill.fgColor.rgb if gt_cell.fill.patternType else None),
            "00D9D9D9")
        loc_rules = [r.formula[0] for rng in ws.conditional_formatting
                     for r in rng.rules
                     if f"{gt_cell.column_letter}$1" in r.formula[0]]
        self.assertTrue(any("COUNTIFS('Backend Data'!$A:$A" in f for f in loc_rules))
        # The grey rule only fires for types Customization knows.
        self.assertTrue(any("$A:$A," in f and ">0" in f for f in loc_rules))

    def test_customization_records_observed_standard_fields(self):
        rows = {(t, h) for t, h, *_ in customization_rows([gt(), air()], {})}
        self.assertIn(("Airfare", "Location"), rows)
        self.assertNotIn(("Ground Transportation", "Location"), rows)
        self.assertIn(("Ground Transportation", "vendor"), rows)
        undef = ExpenseRecord("JJJJ0000", REPORT, 10, [
            cell("expenseType", "Expense Type", "Undefined"),
            cell("businessPurpose", "*Purchase Description/Justification (64 char max)", "x"),
        ])
        urows = {(t, h) for t, h, *_ in customization_rows([undef], {})}
        self.assertIn(("Undefined", "Desc./Justification"), urows)

    def test_orig_sheet_backs_blocked_deviation_rules(self):
        write_workbook(self.path, data())
        wb = load_workbook(self.path)
        self.assertEqual(wb["Orig"].sheet_state, "hidden")
        ws, orig = wb[EXPENSES], wb["Orig"]
        headers = [c.value for c in ws[1]]
        self.assertEqual([c.value for c in orig[1]], headers)
        date = headers.index("Date") + 1
        self.assertEqual(orig.cell(row=2, column=date).value, "09/21/2026")
        dev = [r.formula[0] for rng in ws.conditional_formatting
               for r in rng.rules if "Orig!" in r.formula[0]]
        self.assertTrue(dev)
        self.assertTrue(all("MATCH(" in f for f in dev))
        # The row lookup must search Orig's Expense ID column, wherever the
        # column order put it; a wrong column marks every locked cell red.
        idl = get_column_letter(headers.index("Expense ID") + 1)
        self.assertTrue(all(f"Orig!${idl}:${idl}" in f for f in dev))

    def test_customization_key_column_backs_required_lookup(self):
        write_workbook(self.path, data())
        custom = load_workbook(self.path)["Backend Data"]
        self.assertEqual(custom["G3"].value, "Key")
        self.assertEqual(custom["G4"].value, '=A4&"|"&B4')

    def test_meal_rows_move_to_read_only_sheet(self):
        meal = ExpenseRecord("MMMM0000", REPORT, 9, [
            cell("expenseType", "Expense Type", "Travel Meals - International"),
            cell("transactionDate", "Transaction Date", "07/24/2026", locked=True),
            cell("custom24", "Number of Breakfasts", "1"),
            cell("custom17", "Departure Time (first day of travel)", "5. Whole Day"),
        ])
        d = WorkbookData(REPORT, [gt(), meal])
        write_workbook(self.path, d)
        self.assertEqual(verify_workbook(self.path, d), [])
        wb = load_workbook(self.path)
        ws, ms = wb[EXPENSES], wb["Meals & Per Diem"]
        exp_ids = [r[0].value for r in ws.iter_rows(min_row=2)]
        self.assertNotIn("MMMM0000", exp_ids)
        self.assertEqual(ms["A1"].value, "READ-ONLY (Do not edit)")
        self.assertTrue(ms.protection.sheet)
        mheaders = [c.value for c in ms[3]]
        self.assertIn("Number of Breakfasts", mheaders)
        self.assertIn("Departure Time (first day of travel)", mheaders)
        mrow = {h: v for h, v in zip(mheaders, [c.value for c in ms[4]])}
        self.assertEqual(mrow["Expense Type"], "Travel Meals - International")
        self.assertEqual(mrow["Number of Breakfasts"], "1")
        # Meal rows are invisible to write-back reads.
        from concurrer.xlsxio import read_workbook
        _, rows = read_workbook(self.path)
        self.assertNotIn("MMMM0000", rows)


if __name__ == "__main__":
    unittest.main()
